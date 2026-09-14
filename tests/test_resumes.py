from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from fastapi.testclient import TestClient

from app.core.config import settings


REGISTER_URL = "/api/v1/auth/register"
LOGIN_URL = "/api/v1/auth/login"
RESUMES_URL = "/api/v1/resumes"


def create_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()

    page.insert_text(
        (72, 72),
        "Anish Kumar Verma\nPython FastAPI MySQL Developer",
    )

    pdf_content = document.tobytes()
    document.close()

    return pdf_content


def create_docx() -> bytes:
    document = Document()

    document.add_heading("Anish Kumar Verma", level=1)
    document.add_paragraph(
        "Python, FastAPI, MySQL and REST API development"
    )

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def register_and_login(
    client: TestClient,
    email: str,
) -> dict[str, str]:
    register_response = client.post(
        REGISTER_URL,
        json={
            "full_name": "Anish Kumar Verma",
            "email": email,
            "password": "StrongPass123",
        },
    )

    assert register_response.status_code == 201

    login_response = client.post(
        LOGIN_URL,
        json={
            "email": email,
            "password": "StrongPass123",
        },
    )

    assert login_response.status_code == 200

    access_token = login_response.json()["access_token"]

    return {
        "Authorization": f"Bearer {access_token}"
    }


def upload_pdf(
    client: TestClient,
    headers: dict[str, str],
):
    return client.post(
        RESUMES_URL,
        headers=headers,
        files={
            "file": (
                "anish_resume.pdf",
                create_pdf(),
                "application/pdf",
            )
        },
    )


def test_upload_pdf_resume(
    client: TestClient,
) -> None:
    headers = register_and_login(
        client,
        "anish@example.com",
    )

    response = upload_pdf(client, headers)

    assert response.status_code == 201

    body = response.json()

    assert body["id"] == 1
    assert body["original_filename"] == "anish_resume.pdf"
    assert body["file_type"] == "pdf"
    assert body["file_size"] > 0
    assert body["processing_status"] == "PROCESSED"

    saved_files = list(
        settings.resume_upload_dir.rglob("*.pdf")
    )

    assert len(saved_files) == 1


def test_upload_docx_resume(
    client: TestClient,
) -> None:
    headers = register_and_login(
        client,
        "anish@example.com",
    )

    response = client.post(
        RESUMES_URL,
        headers=headers,
        files={
            "file": (
                "anish_resume.docx",
                create_docx(),
                (
                    "application/vnd.openxmlformats-officedocument"
                    ".wordprocessingml.document"
                ),
            )
        },
    )

    assert response.status_code == 201
    assert response.json()["file_type"] == "docx"
    assert response.json()["processing_status"] == "PROCESSED"


def test_resume_upload_requires_authentication(
    client: TestClient,
) -> None:
    response = client.post(
        RESUMES_URL,
        files={
            "file": (
                "anish_resume.pdf",
                create_pdf(),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 401


def test_reject_unsupported_file_type(
    client: TestClient,
) -> None:
    headers = register_and_login(
        client,
        "anish@example.com",
    )

    response = client.post(
        RESUMES_URL,
        headers=headers,
        files={
            "file": (
                "resume.txt",
                b"Python developer resume",
                "text/plain",
            )
        },
    )

    assert response.status_code == 415
    assert response.json() == {
        "detail": "Only PDF and DOCX resumes are supported"
    }


def test_reject_invalid_pdf(
    client: TestClient,
) -> None:
    headers = register_and_login(
        client,
        "anish@example.com",
    )

    response = client.post(
        RESUMES_URL,
        headers=headers,
        files={
            "file": (
                "fake_resume.pdf",
                b"This is not a real PDF file",
                "application/pdf",
            )
        },
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": "The uploaded file is not a valid PDF"
    }


def test_reject_file_larger_than_limit(
    client: TestClient,
) -> None:
    headers = register_and_login(
        client,
        "anish@example.com",
    )

    oversized_content = (
        b"a" * (settings.max_resume_size_bytes + 1)
    )

    response = client.post(
        RESUMES_URL,
        headers=headers,
        files={
            "file": (
                "large_resume.pdf",
                oversized_content,
                "application/pdf",
            )
        },
    )

    assert response.status_code == 413
    assert response.json() == {
        "detail": "Resume size must not exceed 5 MB"
    }


def test_list_and_get_resume(
    client: TestClient,
) -> None:
    headers = register_and_login(
        client,
        "anish@example.com",
    )

    upload_response = upload_pdf(client, headers)
    resume_id = upload_response.json()["id"]

    list_response = client.get(
        RESUMES_URL,
        headers=headers,
    )

    assert list_response.status_code == 200
    assert len(list_response.json()) == 1

    get_response = client.get(
        f"{RESUMES_URL}/{resume_id}",
        headers=headers,
    )

    assert get_response.status_code == 200
    assert get_response.json()["id"] == resume_id


def test_user_cannot_access_another_users_resume(
    client: TestClient,
) -> None:
    first_user_headers = register_and_login(
        client,
        "first@example.com",
    )

    upload_response = upload_pdf(
        client,
        first_user_headers,
    )

    resume_id = upload_response.json()["id"]

    second_user_headers = register_and_login(
        client,
        "second@example.com",
    )

    get_response = client.get(
        f"{RESUMES_URL}/{resume_id}",
        headers=second_user_headers,
    )

    delete_response = client.delete(
        f"{RESUMES_URL}/{resume_id}",
        headers=second_user_headers,
    )

    assert get_response.status_code == 404
    assert delete_response.status_code == 404


def test_delete_resume_removes_record_and_file(
    client: TestClient,
) -> None:
    headers = register_and_login(
        client,
        "anish@example.com",
    )

    upload_response = upload_pdf(client, headers)
    resume_id = upload_response.json()["id"]

    saved_files = list(
        settings.resume_upload_dir.rglob("*.pdf")
    )

    assert len(saved_files) == 1

    saved_file: Path = saved_files[0]
    assert saved_file.exists()

    delete_response = client.delete(
        f"{RESUMES_URL}/{resume_id}",
        headers=headers,
    )

    assert delete_response.status_code == 204
    assert delete_response.content == b""
    assert not saved_file.exists()

    get_response = client.get(
        f"{RESUMES_URL}/{resume_id}",
        headers=headers,
    )

    assert get_response.status_code == 404