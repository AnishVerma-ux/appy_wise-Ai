from pathlib import Path

import pymupdf
from docx import Document


class ResumeParsingError(Exception):
    """Raised when text cannot be extracted from a resume."""


def _clean_text(text: str) -> str:
    cleaned_lines = []

    for line in text.splitlines():
        cleaned_line = " ".join(line.split())

        if cleaned_line:
            cleaned_lines.append(cleaned_line)

    return "\n".join(cleaned_lines)


def _extract_pdf_text(file_path: Path) -> str:
    text_parts = []

    with pymupdf.open(str(file_path)) as document:
        for page in document:
            page_text = page.get_text("text")

            if page_text:
                text_parts.append(page_text)

    return "\n".join(text_parts)


def _extract_docx_text(file_path: Path) -> str:
    document = Document(str(file_path))
    text_parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # Some resumes use tables for skills, education or experience.
    for table in document.tables:
        for row in table.rows:
            row_values = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_values.append(cell_text)

            if row_values:
                text_parts.append(" | ".join(row_values))

    return "\n".join(text_parts)


def extract_resume_text(file_path: Path, file_type: str) -> str:
    normalized_type = file_type.lower().strip().lstrip(".")

    try:
        if normalized_type == "pdf":
            extracted_text = _extract_pdf_text(file_path)

        elif normalized_type == "docx":
            extracted_text = _extract_docx_text(file_path)

        else:
            raise ResumeParsingError(
                "Only PDF and DOCX resumes are supported"
            )

    except ResumeParsingError:
        raise

    except Exception as exc:
        raise ResumeParsingError(
            "The resume could not be read"
        ) from exc

    cleaned_text = _clean_text(extracted_text)

    if not cleaned_text:
        raise ResumeParsingError(
            "No readable text was found in the resume"
        )

    return cleaned_text